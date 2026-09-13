from __future__ import annotations

import csv
import json
import os
import re
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLORS = {
    "background": "FFFFFFFF",
    "text": "FF111827",
    "subtle": "FF4B5563",
    "header": "FF1F4E78",
    "section": "FF475569",
    "white": "FFFFFFFF",
    "border": "FF94A3B8",
    "border_light": "FFCBD5E1",
    "divider": "FF94A3B8",
    "x": "FFDCEBFA",
    "b": "FFDCFCE7",
    "d": "FFFFEDD5",
    "f": "FFFECACA",
    "c": "FFDCEBFA",
    "r": "FFE9D5FF",
    "mix": "FFFEF3C7",
    "dash": "FFF3F4F6",
}

EXPLANATIONS = {
    "ABB": "A + 2×B",
    "A[K/Q]x": "A + K/Q + ≤9",
    "A[J-T]x": "A + J/T",
    "A[9-2]x": "A + 9…2",
    "BBB": "3×B без A",
    "BBx": "2×B + ≤9",
    "K/Qx": "K/Q high",
    "[J-8]x dis": "J…8 dry",
    "[J-8]x con": "J…8 connected",
    "[7-4]x": "top ≤7",
}

ACTION_NAMES = {
    "check": "check",
    "donk": "donk",
    "bet": "bet",
    "fold": "fold",
    "call": "call",
    "raise": "raise",
}


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _font(size: float = 11, *, bold: bool = False, italic: bool = False,
          color: str = COLORS["text"], display: bool = False) -> Font:
    return Font(
        name="Aptos Display" if display else "Aptos",
        size=size,
        bold=bold,
        italic=italic,
        color=color,
    )


def _side(color: str, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def _all_border(color: str) -> Border:
    side = _side(color)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_cells(ws, cell_range: str, *, fill=None, font=None, alignment=None,
                 border=None, number_format: str | None = None) -> None:
    for row in ws[cell_range]:
        for cell in row:
            if fill is not None:
                cell.fill = copy(fill)
            if font is not None:
                cell.font = copy(font)
            if alignment is not None:
                cell.alignment = copy(alignment)
            if border is not None:
                cell.border = copy(border)
            if number_format is not None:
                cell.number_format = number_format


def _style_base(ws, last_col: int, last_row: int) -> None:
    ws.sheet_view.showGridLines = False
    _style_cells(
        ws,
        f"A1:{get_column_letter(last_col)}{last_row}",
        fill=_fill(COLORS["background"]),
        font=_font(),
        alignment=Alignment(vertical="center"),
    )


def _style_action(cell, value: str) -> None:
    if value == "—":
        color = COLORS["dash"]
    elif "/" in value:
        color = COLORS["mix"]
    elif value == "X":
        color = COLORS["x"]
    elif value == "B":
        color = COLORS["b"]
    elif value == "D":
        color = COLORS["d"]
    elif value == "F":
        color = COLORS["f"]
    elif value in {"C", "BDFD"}:
        color = COLORS["c"]
    elif value == "R":
        color = COLORS["r"]
    else:
        raise RuntimeError(f"Unsupported workbook action label: {value}")
    cell.fill = _fill(color)
    cell.font = _font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _all_border(COLORS["background"])


def _set_bottom_border(ws, row: int, last_col: int, *,
                       color: str = COLORS["divider"],
                       mirror_top: bool = False) -> None:
    line = _side(color, "medium")
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col)
        cell.border = Border(
            left=copy(cell.border.left),
            right=copy(cell.border.right),
            top=_side(color) if mirror_top else copy(cell.border.top),
            bottom=copy(line),
        )


def _set_right_border(ws, col: int, first_row: int, last_row: int) -> None:
    right = _side(COLORS["divider"], "medium")
    for row in range(first_row, last_row + 1):
        cell = ws.cell(row, col)
        cell.border = Border(
            left=copy(cell.border.left),
            right=right,
            top=copy(cell.border.top),
            bottom=copy(cell.border.bottom),
        )


def _frequency_text(values: dict, labels: dict) -> str:
    return "; ".join(
        f"{labels[action]} {100 * value:.1f}%"
        for action, value in values.items()
    )


def _branch_rule(labels: dict) -> str:
    explanation = ", ".join(
        f"{code} = {ACTION_NAMES[action]}" for action, code in labels.items()
    )
    bdfd = (
        " BDFD = fold без BDFD, call с BDFD."
        if "fold" in labels and "call" in labels else ""
    )
    return (
        f"{explanation}. Код через / — ровно 50/50.{bdfd} "
        "Bet и donk = 50% банка."
    )


def _load_model(analysis_root: Path, branches: list[dict]) -> tuple[dict, dict]:
    model_root = analysis_root / "strategy-10"
    manifest = json.loads((model_root / "manifest.json").read_text(encoding="utf-8-sig"))
    loaded = {}
    for branch in branches:
        csv_path = model_root / branch["id"] / "strategy.csv"
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            records = {row["hand_category"]: row for row in csv.DictReader(handle)}
        info = next(item for item in manifest["branches"] if item["branch"] == branch["id"])
        loaded[branch["id"]] = {"records": records, "info": info}
    return manifest, loaded


def _add_summary(wb: Workbook, study: dict, branches: list[dict],
                 manifest: dict, loaded: dict) -> None:
    ws = wb.active
    ws.title = "Summary"
    _style_base(ws, 10, 23)
    ws.merge_cells("A2:J2")
    ws.merge_cells("A3:J3")
    ws["A2"] = study["strategy_title"]
    ws["A3"] = (
        "12 категорий рук × 10 итоговых категорий флопов. "
        "Чистое действие, строгий микс 50/50 или BDFD."
    )
    _style_cells(
        ws, "A2:J2", fill=_fill(COLORS["background"]),
        font=_font(20, bold=True, display=True),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    _set_bottom_border(ws, 2, 10, color=COLORS["header"])
    _style_cells(
        ws, "A3:J3", fill=_fill(COLORS["background"]),
        font=_font(11, color=COLORS["subtle"]),
    )
    headers = [
        "Ситуация", "Игрок", "Reach узла", "Solver frequencies", "Стратегия",
        "Loss в узле, bb", "Loss от root, bb", "BDFD ячеек",
        "Комбо reach > 0", "Ячеек",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    _style_cells(
        ws, "A5:J5", fill=_fill(COLORS["header"]),
        font=_font(10, bold=True, color=COLORS["white"]),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_all_border(COLORS["white"]),
    )
    for row_index, branch in enumerate(branches, 6):
        info = loaded[branch["id"]]["info"]
        summary = info["summary"]
        values = [
            branch["strategy_title"],
            branch["acting_player"],
            summary["node_reach_fraction"],
            _frequency_text(summary["solver_reach_weighted_frequencies"], info["labels"]),
            _frequency_text(summary["simplified_reach_weighted_frequencies"], info["labels"]),
            summary["mean_local_loss_bb"],
            summary["root_loss_bb"],
            summary["bdfd_cells"],
            summary["active_rows"],
            summary["populated_cells"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    last_branch_row = 5 + len(branches)
    _style_cells(
        ws, f"A6:J{last_branch_row}",
        font=_font(10),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_all_border(COLORS["border"]),
    )
    _style_cells(ws, f"C6:C{last_branch_row}", number_format="0.0%")
    _style_cells(ws, f"F6:G{last_branch_row}", number_format="0.000000")
    _style_cells(ws, f"H6:J{last_branch_row}", number_format="0")

    for col, value in enumerate(["Код", "Действие", "Цвет", "Правило"], 1):
        ws.cell(13, col, value)
    ws["F13"] = "Категории флопов"
    _style_cells(
        ws, "A13:D13", fill=_fill(COLORS["section"]),
        font=_font(10, bold=True, color=COLORS["white"]),
        alignment=Alignment(horizontal="center", vertical="center"),
        border=_all_border(COLORS["border_light"]),
    )
    _style_cells(
        ws, "F13:J13", fill=_fill(COLORS["section"]),
        font=_font(10, bold=True, color=COLORS["white"]),
        alignment=Alignment(horizontal="center", vertical="center"),
        border=_all_border(COLORS["border_light"]),
    )
    legend = [
        ["X", "Check", "синий", None],
        ["B", "Bet 50%", "зелёный", None],
        ["D", "Donk 50%", "оранжевый", None],
        ["F", "Fold", "красный", None],
        ["C", "Call", "голубой", None],
        ["R", "Raise", "фиолетовый", None],
        ["BDFD", "Call только с BDFD", "голубой", "без BDFD — fold"],
        ["X/B, F/C и т. п.", "Два действия", "жёлтый", "строго 50/50"],
        ["—", "Нет категории", "серый", "в диапазоне/классе отсутствует"],
    ]
    for row_index, values in enumerate(legend, 14):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    _style_cells(
        ws, "A14:D22", font=_font(10),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_all_border(COLORS["border_light"]),
    )
    notes = [
        "10 итоговых категорий; B13 используется только для однозначного разбиения",
        "Каждый реальный флоп внутри итоговой категории получает одинаковый вес.",
        "12 категорий рук; OESD и Gutshot имеют приоритет над made hand.",
        "Все значения через / означают только микс 50/50.",
        "EV аудирует потери; только для BDFD служит safety-фильтром.",
    ]
    for row_index, note in enumerate(notes, 14):
        ws.merge_cells(start_row=row_index, start_column=6, end_row=row_index, end_column=10)
        ws.cell(row_index, 6, note)
    _style_cells(
        ws, "F14:J18", font=_font(10),
        alignment=Alignment(vertical="center", wrap_text=True),
    )
    widths = [27, 11, 12, 23, 23, 16, 16, 16, 13, 13]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[5].height = 36
    for row in range(14, 23):
        ws.row_dimensions[row].height = 26
    ws.freeze_panes = "A6"


def _add_strategy_sheet(wb: Workbook, branch: dict, loaded: dict,
                        hand_order: list[str], flop_groups: list[dict]) -> None:
    info = loaded[branch["id"]]["info"]
    ws = wb.create_sheet(branch["strategy_sheet"])
    last_col = 1 + len(flop_groups)
    last_row = 6 + len(hand_order)
    _style_base(ws, last_col, last_row)
    last_letter = get_column_letter(last_col)
    ws.merge_cells(f"A2:{last_letter}2")
    ws.merge_cells(f"A3:{last_letter}3")
    ws["A2"] = branch["strategy_title"]
    ws["A3"] = _branch_rule(info["labels"])
    _style_cells(
        ws, f"A2:{last_letter}2", font=_font(18, bold=True, display=True),
        alignment=Alignment(vertical="center"),
    )
    _set_bottom_border(ws, 2, last_col, color=COLORS["header"])
    _style_cells(
        ws, f"A3:{last_letter}3",
        font=_font(10, color=COLORS["subtle"]),
    )
    headers = ["Категория руки"] + [
        f"{group['name']} ({group['count']})" for group in flop_groups
    ]
    explanations = [None] + [EXPLANATIONS[group["name"]] for group in flop_groups]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    for col, value in enumerate(explanations, 1):
        ws.cell(6, col, value)
    _style_cells(
        ws, f"A5:{last_letter}5", fill=_fill(COLORS["header"]),
        font=_font(9, bold=True, color=COLORS["white"]),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_all_border(COLORS["white"]),
    )
    _style_cells(
        ws, f"A6:{last_letter}6", fill=_fill(COLORS["header"]),
        font=_font(8, italic=True, color=COLORS["white"]),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_all_border(COLORS["header"]),
    )
    for row_index, hand in enumerate(hand_order, 7):
        record = loaded[branch["id"]]["records"].get(hand)
        if record is None:
            raise RuntimeError(f"{branch['id']}: missing hand row {hand}")
        ws.cell(row_index, 1, hand)
        for col, group in enumerate(flop_groups, 2):
            value = record[group["name"]]
            ws.cell(row_index, col, value)
            _style_action(ws.cell(row_index, col), value)
    _style_cells(
        ws, f"A7:A{last_row}", font=_font(10, bold=True),
        alignment=Alignment(vertical="center"),
        border=_all_border(COLORS["border"]),
    )
    for row in (14, 16, 17):
        _set_bottom_border(ws, row, last_col, mirror_top=True)
    for offset in (2, 4, 6, 7, 8, 9, 10):
        _set_right_border(ws, 1 + offset, 5, last_row)
    ws.column_dimensions["A"].width = 23
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.row_dimensions[5].height = 34
    ws.row_dimensions[6].height = 30
    for row in range(7, last_row + 1):
        ws.row_dimensions[row].height = 24
    ws.freeze_panes = "B7"


def _add_method(wb: Workbook, study: dict, flop_groups: list[dict]) -> None:
    ws = wb.create_sheet("Method")
    last_row = 21 + len(flop_groups)
    _style_base(ws, 4, last_row)
    ws.merge_cells("A2:D2")
    ws["A2"] = "Методика"
    _style_cells(
        ws, "A2:D2", font=_font(20, bold=True, display=True),
        alignment=Alignment(vertical="center"),
    )
    _set_bottom_border(ws, 2, 4, color=COLORS["section"])
    ws["A4"] = "Параметр"
    ws["B4"] = "Правило"
    _style_cells(
        ws, "A4:D4", fill=_fill(COLORS["section"]),
        font=_font(10, bold=True, color=COLORS["white"]),
        border=_all_border(COLORS["border_light"]),
    )
    method_rows = [
        ["Флопы", "286 canonical unpaired rainbow flops, 10 итоговых взаимоисключающих категорий"],
        ["Диапазоны", study["strategy_range_note"]],
        ["Частоты", "Равное среднее по concrete combo на флопе, затем равное среднее по реальным флопам категории"],
        ["Чистое действие", "Наибольшая средняя частота строго >65%"],
        ["Микс", "Иначе два наиболее частых действия, строго 50/50; первым указано более частое действие солвера"],
        ["Вес диапазона", "Combo с reach_probability >0 участвует независимо от величины reach"],
        ["EV-аудит", "Reach-weighted local regret против solved opponent; не adaptive exploitability. Обычное действие выбирают частоты; EV только отклоняет невыгодный BDFD-rule"],
        ["Категории рук", "8 made-категорий, OESD, Gutshot, 2 overcards + BDFD и Air"],
        ["Карманные пары", "Underpair: между top и middle; Weak pair: между middle и low; Low pocket pair: ниже low"],
        ["Приоритет классификации", "OESD > Gutshot > made hand > 2 overcards + BDFD > Air"],
        ["BDFD в ячейке", "Без BDFD fold >65%; с BDFD суммарный continue >65%, который упрощается до call; затем правило проходит EV safety audit"],
        ["Категории флопов", "10 итоговых категорий; B13 используется только для однозначного разбиения"],
        ["Роль исходной B13", "Только однозначное внутреннее разбиение 286 флопов на 10 итоговых групп"],
        ["JT9", "Исходная B13 = [J-8]x con; итоговая категория = [J-8]x con"],
    ]
    for row_index, values in enumerate(method_rows, 5):
        ws.cell(row_index, 1, values[0])
        ws.cell(row_index, 2, values[1])
    _style_cells(
        ws, "A5:B18", font=_font(10),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_all_border(COLORS["border_light"]),
    )
    ws.append([])
    ws.append([])
    header_row = 21
    for col, value in enumerate(
        ["Категория флопов", "Входящие B13 / определение", "Количество", "Пояснение"], 1
    ):
        ws.cell(header_row, col, value)
    _style_cells(
        ws, f"A{header_row}:D{header_row}", fill=_fill(COLORS["section"]),
        font=_font(10, bold=True, color=COLORS["white"]),
        border=_all_border(COLORS["white"]),
    )
    explanations = {
        "ABB": "A + две broadway-карты",
        "A[K/Q]x": "A + K/Q + карта 9 и ниже",
        "A[J-T]x": "A с J/T",
        "A[9-2]x": "A со средней картой 9 или ниже",
        "BBB": "Без A; три broadway-карты",
        "BBx": "Без A; две broadway + карта 9 и ниже; JT9 исключён",
        "K/Qx": "Старшая K/Q, но не BBx",
        "[J-8]x dis": "Старшая J/T/9/8; две младшие не соседние",
        "[J-8]x con": "Старшая J/T/9/8; две младшие соседние; включает JT9",
        "[7-4]x": "Старшая карта 7 или ниже",
    }
    for row_index, group in enumerate(flop_groups, 22):
        members = " + ".join(group["members"])
        values = [group["name"], members, group["count"], explanations[group["name"]]]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    _style_cells(
        ws, f"A22:D{last_row}", font=_font(9),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_all_border(COLORS["border_light"]),
    )
    _style_cells(ws, f"C22:C{last_row}", number_format="0")
    for col, width in zip(("A", "B", "C", "D"), (24, 52, 12, 54)):
        ws.column_dimensions[col].width = width
    for row in range(5, 19):
        ws.row_dimensions[row].height = 42
    for row in range(22, last_row + 1):
        ws.row_dimensions[row].height = 30


def _save_deterministic(wb: Workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.creator = "PromSoftService"
    wb.properties.lastModifiedBy = "PromSoftService"
    wb.properties.created = datetime(2000, 1, 1)
    wb.properties.modified = datetime(2000, 1, 1)
    with tempfile.TemporaryDirectory(prefix=".flop-workbook-", dir=output_path.parent) as temp_dir:
        raw_path = Path(temp_dir) / "raw.xlsx"
        packed_path = Path(temp_dir) / "packed.xlsx"
        wb.save(raw_path)
        with zipfile.ZipFile(raw_path, "r") as source, zipfile.ZipFile(
            packed_path, "w", compression=zipfile.ZIP_STORED
        ) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_STORED
                info.create_system = 0
                info.external_attr = 0
                data = source.read(name)
                if name == "docProps/core.xml":
                    data = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                        data,
                    )
                target.writestr(info, data)
        os.replace(packed_path, output_path)


def _validate_saved(output_path: Path, branches: list[dict], loaded: dict,
                    hand_order: list[str], flop_groups: list[dict]) -> dict:
    with zipfile.ZipFile(output_path, "r") as archive:
        bad = archive.testzip()
        if bad is not None:
            raise RuntimeError(f"Corrupt XLSX member: {bad}")
    wb = load_workbook(output_path, data_only=False)
    expected_sheets = ["Summary"] + [branch["strategy_sheet"] for branch in branches] + ["Method"]
    if wb.sheetnames != expected_sheets:
        raise RuntimeError(f"Workbook sheet mismatch: {wb.sheetnames} != {expected_sheets}")
    formula_errors = []
    checked = 0
    for branch in branches:
        ws = wb[branch["strategy_sheet"]]
        if ws.freeze_panes != "B7" or ws.max_row != 18 or ws.max_column != 11:
            raise RuntimeError(f"{ws.title}: layout mismatch")
        for row_index, hand in enumerate(hand_order, 7):
            record = loaded[branch["id"]]["records"][hand]
            if ws.cell(row_index, 1).value != hand:
                raise RuntimeError(f"{ws.title}: hand row mismatch at {row_index}")
            for col, group in enumerate(flop_groups, 2):
                expected = record[group["name"]]
                actual = ws.cell(row_index, col).value
                if actual != expected:
                    raise RuntimeError(
                        f"{ws.title}!{get_column_letter(col)}{row_index}: {actual} != {expected}"
                    )
                checked += int(actual != "—")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    token in cell.value
                    for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")
                ):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if formula_errors:
        raise RuntimeError(f"Workbook formula errors: {formula_errors}")
    return {
        "path": str(output_path),
        "sheets": len(wb.sheetnames),
        "populated_strategy_cells": checked,
        "formula_errors": 0,
    }


def build_flop_workbook(repo: Path, study_directory: str,
                        hand_order: list[str], flop_groups: list[dict]) -> dict:
    study_path = repo / "studies" / study_directory / "study.json"
    study = json.loads(study_path.read_text(encoding="utf-8-sig"))
    branches = study["branches"]
    for branch in branches:
        if not branch.get("strategy_sheet") or not branch.get("strategy_title"):
            raise RuntimeError(f"{study_directory}/{branch['id']}: workbook metadata missing")
    analysis_root = repo / "datasets" / study_directory / "analysis"
    manifest, loaded = _load_model(analysis_root, branches)
    wb = Workbook()
    _add_summary(wb, study, branches, manifest, loaded)
    for branch in branches:
        _add_strategy_sheet(wb, branch, loaded, hand_order, flop_groups)
    _add_method(wb, study, flop_groups)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    output_path = analysis_root / f"{study['study_id']}_flop_strategy.xlsx"
    _save_deterministic(wb, output_path)
    return _validate_saved(output_path, branches, loaded, hand_order, flop_groups)
